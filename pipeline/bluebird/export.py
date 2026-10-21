"""공유용 내보내기.

- xlsx: 아이디어 전체 + 선별 재료 → 엑셀 1파일. share 열(가능/보류) = 레지스트리 share → core.source.share_ok
- master_xlsx: 정제 DB master 표 → 엑셀 1파일(시트 master)
- provenance: docs/DATA.md 소스 표. 원천 = 정제 DB(master·raw) + 레지스트리(TOML)
- 읽기 전용. 공개 승인·깔때기 무관, core 전체
"""
from __future__ import annotations

import datetime as dt
from collections import defaultdict
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from . import db
from .sources import load_sources

IDEA_COLS = ("id", "source_id", "source_name", "license", "share", "record_kind", "contest", "host_org", "audience",
             "year", "award", "title", "body", "category", "used_data", "source_url", "exclude_reasons", "canonical_id")

IDEA_SQL = """
SELECT i.id, s.id, s.name, s.license, s.share_ok, s.record_kind, coalesce(nullif(c.name, ''), s.name), c.host_org, c.audience,
       i.year, i.award, i.title, i.body, i.category, array_to_string(i.used_data, '; '), i.source_url,
       (SELECT string_agg(x.reason, ', ' ORDER BY x.reason) FROM core.idea_exclusion x WHERE x.idea_id = i.id),
       (SELECT d.canonical_id FROM core.idea_duplicate d WHERE d.idea_id = i.id)
  FROM core.idea i
  JOIN core.source s ON s.id = i.source_id
  JOIN core.contest c ON c.id = i.contest_id
 WHERE i.retired_at IS NULL AND i.withheld_at IS NULL
 ORDER BY s.id, i.year NULLS LAST, i.id
"""


SOURCE_SQL = """
SELECT s.id, s.name, s.license, s.url, s.record_kind, s.exclude_reason,
       coalesce(i.n, 0), i.y0, i.y1, coalesce(i.with_body, 0), r.first_at, r.last_at, s.share_ok
  FROM core.source s
  LEFT JOIN (SELECT source_id, count(*) n, min(year) y0, max(year) y1, count(body) with_body
               FROM core.idea WHERE retired_at IS NULL GROUP BY 1) i ON i.source_id = s.id
  LEFT JOIN (SELECT source_id, min(started_at)::date first_at, max(started_at)::date last_at
               FROM core.ingest_run WHERE status = 'ok' GROUP BY 1) r ON r.source_id = s.id
 WHERE coalesce(i.n, 0) > 0   -- 퇴역된 소스(유효 0건) 제외
 ORDER BY i.n DESC
"""


MASTER_COLS = ("api", "key", "title", "body", "kind", "field", "program", "award", "year", "url", "exclude_reasons",
               "label", "item_type", "goal", "solution", "features", "domain", "evidence", "verified")

# master 칸 순서 그대로. 목록(exclude_reasons·features·domain) = 쉼표 이음
MASTER_SQL = """
SELECT api, key, title, body, kind, field, program, award, year, url, array_to_string(exclude_reasons, ', '),
       label, item_type, goal, solution, array_to_string(features, ', '), array_to_string(domain, ', '), evidence, verified
  FROM master
 ORDER BY api, year NULLS LAST, key
"""

MASTER_WIDTHS = {"api": 24, "key": 18, "title": 40, "body": 60, "kind": 10, "field": 14, "program": 28, "award": 12,
                 "year": 7, "url": 30, "exclude_reasons": 14, "label": 10, "item_type": 10, "goal": 40, "solution": 40,
                 "features": 30, "domain": 14, "evidence": 30, "verified": 8}
XLSX_CELL_MAX = 32767  # 엑셀 칸 글자 수 한도


def _text_cell(ws, value):
    """문자열은 항상 텍스트 칸. '='로 시작하는 본문이 수식으로 저장되면 엑셀이 '복구' 경고를 띄운다."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:XLSX_CELL_MAX]
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def master_xlsx(*, dsn: str, out: Path) -> dict:
    """master 표 전체 → 시트 'master' 1개(MASTER_COLS). docs/MASTER.md = 필드 설명."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("master")
    for idx, name in enumerate(MASTER_COLS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = MASTER_WIDTHS[name]
    ws.freeze_panes = "A2"
    ws.append(list(MASTER_COLS))
    n = 0
    with db.connect(dsn) as conn, conn.cursor(name="master_xlsx") as cur:
        cur.itersize = 5000
        cur.execute(MASTER_SQL)
        for row in cur:
            ws.append([_text_cell(ws, v) for v in row])
            n += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(out)
    return {"path": str(out), "rows": n, "cols": len(MASTER_COLS)}


def _autosize(ws, widths: dict[str, int]):
    for idx, cell in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(cell.value, 14)
    ws.freeze_panes = "A2"


def _share(ok: bool) -> str:
    return "가능" if ok else "보류"


def xlsx(*, dsn: str, out: Path) -> dict:
    with db.connect(dsn) as conn:
        sources = conn.execute(SOURCE_SQL).fetchall()
        ideas = conn.execute(IDEA_SQL).fetchall()
        reasons = conn.execute("SELECT code, label, basis, note FROM core.exclusion_reason ORDER BY code").fetchall()
        reason_counts = dict(conn.execute("SELECT reason, count(*) FROM core.idea_exclusion GROUP BY 1").fetchall())
        contests = conn.execute(
            """SELECT c.id, c.source_id, c.name, c.host_org, c.year, c.audience, count(i.id)
                 FROM core.contest c LEFT JOIN core.idea i ON i.contest_id = c.id AND i.retired_at IS NULL
                GROUP BY c.id ORDER BY count(i.id) DESC""").fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "caution"
    today = dt.date.today().isoformat()
    n_ideas = len(ideas)
    n_share = sum(1 for r in ideas if r[4])
    lines = [
        ("파랑새 아이디어 데이터 공유본", "", ""),
        ("생성일", today, ""),
        ("아이디어 건수", n_ideas, ""),
        ("소스 수", len(sources), ""),
        ("share=가능", f"{n_share}건 / 소스 {sum(1 for r in sources if r[12])}개", ""),
        ("share=보류", f"{n_ideas - n_share}건 / 소스 {sum(1 for r in sources if not r[12])}개", ""),
        ("", "", ""),
        ("주의", "수상자·제안자 이름: 적재 시 제거. 본문 내 잔존 이름 = 원 기관 공개본 그대로", ""),
        ("주의", "share='보류' 행: 원 기관 약관 확인 전 외부 재배포 금지", ""),
        ("주의", "exclude_reasons: '뺀다면 이 이유' 후보. 필터링 안 함. 빈 값 = 제외 근거 없음", ""),
        ("", "", ""),
        ("시트", "내용", ""),
        ("ideas", "아이디어 1건 = 1행", ""),
        ("sources", "소스별 출처 URL·라이선스·건수·연도·수신일", ""),
        ("contests", "대회별 건수·대상자", ""),
        ("exclusion_reason", "ideas.exclude_reasons 코드 뜻·건수", ""),
        ("", "", ""),
        ("ideas 시트 필드", "설명", "예시·값"),
        ("id", "아이디어 고유 ID. ID-연도-해시10. 재적재 시 동일 건 = 동일 id", "ID-2021-3f9a1c2b7e"),
        ("source_id", "소스 식별자. sources.source_id 참조", "gov_startup_prelim_awards"),
        ("source_name", "소스 이름(기관_데이터셋명)", "범정부 공공데이터 활용 창업경진대회 예선기관 수상작"),
        ("license", "원 기관 이용 조건. 공공데이터포털 표기 또는 기관 약관 원문", "이용허락범위 제한 없음 / 공공누리 제1유형"),
        ("share", "외부 공유 가능 여부. 소스별로 사람이 확인해 레지스트리(share)에 적은 값. 보류 = 원 기관 이용조건 확인 전", "가능 / 보류"),
        ("record_kind", "기록 종류. contest_award=수상작 / contest_submission=출품작(수상 여부 불명) / public_proposal=국민·주민 제안(참여예산 포함) / improvement_proposal=직원 개선제안", "contest_award"),
        ("contest", "대회·공모 이름. 원본에 없으면 소스 이름", "제42회 전국학생과학발명품경진대회"),
        ("host_org", "주관·제공 기관", "국립중앙과학관"),
        ("audience", "대회 대상자. student=초·중·고 학생 대회(이름에 학생·청소년·과학전람회 등) / general=일반(사람 확정) / unknown=미확인. '대학생'은 student 아님", "student / general / unknown"),
        ("year", "수상·제안 연도. 원본에 없으면 빈 값", "2021"),
        ("award", "상 이름·등급. 원본 표기 그대로. 출품작·제안은 대개 빈 값", "대상 / 우수상 / 장려상 / 특상"),
        ("title", "아이디어 제목. 사람 이름 포함 팀명은 제거", "저상버스 도착 알림 서비스"),
        ("body", "설명·내용 본문. 본문 없는 소스(참여예산 대부분, 과학관 2015~)는 빈 값", "(긴 글)"),
        ("category", "원본 분류·분야. 소스별 체계 상이(통일 안 함)", "환경/에너지 / 정보통신"),
        ("used_data", "활용 공공데이터(데이터 활용 경진대회만). ; 구분", "버스정보시스템; 상가업소정보"),
        ("source_url", "원본 위치. 건별 상세 URL, 없으면 데이터셋 페이지", "https://www.data.go.kr/data/15146089/fileData.do"),
        ("exclude_reasons", "제외 후보 이유 코드. 복수 = , 구분. 뜻은 exclusion_reason 시트. 받는 쪽이 골라 제외", "student_contest, public_proposal"),
        ("canonical_id", "같은 아이디어의 다른 기록이면 묶음 대표 id(본문 같음 = 재제출·중복 등록 / 같은 대회·연도·제목 = 출품작·수상작·KIPRIS 목록). 대표·단독 = 빈 값", "ID-2015-3df3a02be2"),
        ("", "", ""),
        ("sources 시트 필드", "설명", ""),
        ("source_id / name / share / license / url", "ideas 시트와 동일. url = 데이터셋 페이지", ""),
        ("record_kind", "소스 단위 기록 종류", ""),
        ("exclude_reason", "소스 단위 편집 판단. certified_product=상품화 후 인증(굿디자인) / internal_activity=기관 내부 활동(품질분임조 등) / teaching_case=수업 사례. 없으면 빈 값", ""),
        ("ideas / year_min / year_max / with_body", "유효 건수 / 연도 범위 / 본문 있는 건수", ""),
        ("first_ingest / last_ingest", "최초·최종 적재일", ""),
        ("", "", ""),
        ("contests 시트 필드", "설명", ""),
        ("contest_id / source_id / name / host_org / year", "대회 식별자·소스·이름·주관·연도. 소스 다르면 별도 행", ""),
        ("audience / ideas", "대상자(ideas 시트와 동일) / 대회별 건수", ""),
        ("", "", ""),
        ("exclusion_reason 시트 필드", "설명", ""),
        ("code / label / note", "이유 코드 / 한글 뜻 / 판단 기준", ""),
        ("basis", "근거 위치. source=소스 전체 / contest=대회 단위 / content=본문·제목 내용(중복 기록·내용 부족)", ""),
        ("ideas", "해당 이유가 붙은 아이디어 건수", ""),
    ]
    for row in lines:
        ws.append(list(row))
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 110
    ws.column_dimensions["C"].width = 48

    ws = wb.create_sheet("ideas")
    ws.append(list(IDEA_COLS))
    for r in ideas:
        ws.append([r[0], r[1], r[2], r[3], _share(r[4]), r[5], r[6], r[7], r[8], r[9], r[10], r[11], r[12], r[13], r[14],
                   r[15], r[16] or "", r[17] or ""])
    _autosize(ws, {"id": 22, "source_id": 30, "source_name": 40, "license": 30, "share": 6, "record_kind": 18,
                   "contest": 30, "host_org": 20, "audience": 9, "year": 6, "award": 12, "title": 50, "body": 80,
                   "category": 14, "used_data": 30, "source_url": 40, "exclude_reasons": 30, "canonical_id": 22})

    ws = wb.create_sheet("sources")
    ws.append(["source_id", "name", "share", "license", "url", "record_kind", "exclude_reason", "ideas",
               "year_min", "year_max", "with_body", "first_ingest", "last_ingest"])
    for r in sources:
        ws.append([r[0], r[1], _share(r[12]), r[2], r[3], r[4], r[5] or "", r[6], r[7], r[8],
                   r[9], str(r[10] or ""), str(r[11] or "")])
    _autosize(ws, {"source_id": 36, "name": 50, "license": 40, "url": 50, "record_kind": 20})

    ws = wb.create_sheet("contests")
    ws.append(["contest_id", "source_id", "name", "host_org", "year", "audience", "ideas"])
    for r in contests:
        ws.append(list(r))
    _autosize(ws, {"contest_id": 20, "source_id": 36, "name": 50, "host_org": 24})

    ws = wb.create_sheet("exclusion_reason")
    ws.append(["code", "label", "basis", "note", "ideas"])
    for code, label, basis, note in reasons:
        ws.append([code, label, basis, note, reason_counts.get(code, 0)])
    _autosize(ws, {"code": 22, "label": 28, "basis": 9, "note": 80})

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"path": str(out), "ideas": n_ideas, "sources": len(sources), "share_ok": n_share}


# ---------------------------------------------------------------- docs/DATA.md

PLATFORM = {"www.data.go.kr": "공공데이터포털", "data.seoul.go.kr": "서울 열린데이터광장",
            "idea.seoul.go.kr": "서울 열린데이터광장", "openapi.gg.go.kr": "경기데이터드림",
            "www.science.go.kr": "국립중앙과학관 누리집", "award.kidp.or.kr": "한국디자인진흥원 누리집",
            "plus.kipris.or.kr": "KIPRIS Plus"}
METHOD = {None: "파일", "data_go_csv": "파일", "science_archive": "크롤링", "kidp_awards": "크롤링"}   # 그 외 = API

MASTER_GROUP_SQL = """
SELECT api, kind, field, program, count(*) FILTER (WHERE u), count(*) FILTER (WHERE u AND coalesce(body, '') <> ''),
       count(*) FILTER (WHERE u AND award IS NOT NULL), min(year) FILTER (WHERE u), max(year) FILTER (WHERE u),
       count(*)
  FROM (SELECT *, NOT 'duplicate' = ANY(exclude_reasons) AS u FROM master) m
 GROUP BY 1, 2, 3, 4
"""


def _years(y0, y1) -> str:
    return f"{y0}~{y1}" if y0 and y1 and y0 != y1 else str(y0 or y1 or "")


def _pick(f, a, b):
    """None 무시하고 f(a, b)."""
    return b if a is None else a if b is None else f(a, b)


def provenance(*, dsn: str, configs: list[Path], out: Path) -> dict:
    """docs/DATA.md — 수집처 → 종류·분야 → 대회·제도별 행 수(중복 제외). 분류 = master.kind·field·program(refine)."""
    specs = {}
    for cfg in configs:
        for s in load_sources(cfg, Path("/nonexistent")):
            specs[s.id] = s
    with db.connect(dsn) as conn:
        groups = conn.execute(MASTER_GROUP_SQL).fetchall()
        raw = dict(conn.execute("SELECT source, count(*) FROM raw GROUP BY 1").fetchall())
        basis = conn.execute("SELECT max(fetched)::date FROM raw").fetchone()[0]
    detail = {s.removesuffix(".detail"): n for s, n in raw.items() if s.endswith(".detail")}

    def place(sid: str) -> tuple[str, str]:
        spec = specs.get(sid)
        host = urlsplit(spec.url).hostname if spec else ""
        return PLATFORM.get(host, host or "미상"), METHOD.get(spec.collector, "API") if spec else "미상"

    # (종류, 분야, 대회·제도) → [행, 본문, 수상, 연도 최소, 연도 최대, 수집처 집합]
    prog: dict[tuple[str, str, str], list] = {}
    per_source: dict[str, int] = defaultdict(int)
    total = dup = 0
    for sid, kind, field, label, n, body, award, y0, y1, all_n in groups:
        total += all_n
        dup += all_n - n
        per_source[sid] += n
        if not n:
            continue
        if kind is None:   # classify 전 행
            spec = specs.get(sid)
            kind, field, label = "미분류", "미분류", spec.name if spec else sid
        p = prog.setdefault((kind, field, label), [0, 0, 0, None, None, set()])
        p[0] += n
        p[1] += body
        p[2] += award
        p[3], p[4] = _pick(min, p[3], y0), _pick(max, p[4], y1)
        p[5].add(place(sid)[0])
    uniq = total - dup

    platforms: dict[str, list] = {}   # 수집처 → [방식별 데이터셋 수, 상세 페이지 크롤링 수, 행]
    for sid, n in per_source.items():
        name, method = place(sid)
        p = platforms.setdefault(name, [defaultdict(int), 0, 0])
        p[0][method] += 1
        p[1] += sid in detail
        p[2] += n
    list_n = sum(n for s, n in raw.items() if not s.endswith(".detail"))
    md = ["# 데이터", "",
          f"- 기준: {basis} (raw 수집 시각 최댓값)",
          f"- 원본(raw) {list_n + sum(detail.values()):,}행 = 목록 {list_n:,} + 상세 페이지 본문 {sum(detail.values()):,}",
          f"- 정제(master) {total:,}행 = 중복 제외 {uniq:,} + 중복 {dup:,}",
          "- 아래 행 수 = 중복 제외", "",
          f"## 수집처 {len(platforms)}곳 · 데이터셋 {len(per_source)}개", "",
          "| 수집처 | 데이터셋(방식별) | 상세 페이지 크롤링 | 행 |", "|---|---|---:|---:|"]
    for name, (methods, crawled, n) in sorted(platforms.items(), key=lambda x: -x[1][2]):
        sets = " · ".join(f"{m} {c}" for m, c in sorted(methods.items(), key=lambda x: -x[1]))
        md.append(f"| {name} | {sets} | {crawled} | {n:,} |")

    fields: dict[tuple[str, str], list] = {}
    for (kind, field, _label), p in prog.items():
        f = fields.setdefault((kind, field), [0, 0, 0, 0])
        f[0] += 1
        f[1] += p[0]
        f[2] += p[1]
        f[3] += p[2]
    order = {k: i for i, k in enumerate(("공모전", "시민 제안", "직원 제안", "미분류"))}
    keys = sorted(fields, key=lambda k: (order.get(k[0], 9), -fields[k][1]))
    md += ["", f"## 종류·분야 (대회·제도 {len(prog)}개)", "",
           "| 종류 | 분야 | 대회·제도 | 행 | 본문 있음 | 수상 등급 있음 |", "|---|---|---:|---:|---:|---:|"]
    for kind, field in keys:
        c, n, body, award = fields[(kind, field)]
        md.append(f"| {kind} | {field} | {c} | {n:,} | {body:,} | {award:,} |")

    md += ["", "## 대회·제도"]
    for kind, field in keys:
        md += ["", f"### {kind} · {field} {fields[(kind, field)][1]:,}행", "",
               "| 대회·제도 | 연도 | 행 | 본문 있음 | 수상 등급 있음 | 수집처 |", "|---|---|---:|---:|---:|---|"]
        rows = sorted(((label, p) for (k, f, label), p in prog.items() if (k, f) == (kind, field)),
                      key=lambda x: -x[1][0])
        for label, (n, body, award, y0, y1, plats) in rows:
            md.append(f"| {label} | {_years(y0, y1)} | {n:,} | {body:,} | {award:,} | {', '.join(sorted(plats))} |")

    raw_only = sorted((s, n) for s, n in raw.items() if not s.endswith(".detail") and s not in per_source)
    if raw_only:
        md += ["", "## 원본만(정제 0행)", ""] + [f"- `{s}` {n:,}행" for s, n in raw_only]
    md += ["", "## 조회", "",
           "- PostgreSQL `161.122.152.68:5432` · DB `parangsae` · 사용자 `aix`",
           "- 표: `master`(최종) · `raw`(원본) · `db_log`(작업 기록). 칸 설명 = DB 주석, 안내 = `docs/MASTER.md`", "",
           "```sql", "SELECT * FROM master WHERE exclude_reasons = '{}';   -- 사용 대상", "```", ""]
    out.write_text("\n".join(md), encoding="utf-8")
    return {"path": str(out), "platforms": len(platforms), "programs": len(prog), "rows": uniq}
