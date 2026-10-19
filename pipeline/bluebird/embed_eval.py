"""임베딩 후보 평가(AC1). 사람 라벨 없이 DB만으로 만든다. 결정은 사용자가 한다.

- 표본 E: 대표 기록(대상 중 core.idea_duplicate에 없는 행) 5,000건, 소스 비례 + C 대상 소스 최소 100건
- A 제목→본문: E 중 제목이 E 안에서 유일한 행의 제목(:title)으로 E의 본문(:notitle, 제목을 지운 본문)을 찾는다
- B 같은 사업명 이웃: 같은 제목(공백 제거 10자 이상, 본문은 다름)을 가진 다른 기록이 본문만으로 가까운지.
  계획의 '같은 공모·다른 소스 쌍'은 실측 0쌍이라(same_entry 기록은 본문이 없음) 이 정의로 대신한다
- C 원본 분류: 같은 소스 안 kNN@10의 같은 category 비율(purity) ÷ 우연 기준선 Σp²
- Matryoshka: 3072 벡터 앞부분만 잘라 다시 정규화(추가 호출 없음)
- 벡터는 embed.ensure로 core.idea_embedding에 키·variant별로 캐시 → 재실행 시 호출 0건
"""

from __future__ import annotations

import math
import time
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook

from . import db, embed
from .egress import Egress

SEED = 20261013
E_SIZE = 5000
A_QUERIES = 2000
B_QUERIES = 1000
B_MAX_GROUP = 50
C_MIN_ROWS = 100
REVIEW_QUERIES = 25
_NT = "openai:text-embedding-3-large:3072:notitle"   # input_text는 variant만 본다(표본 선정용)
DIMS = {"openai:text-embedding-3-large:3072": (3072, 1536, 1024),
        "gemini:gemini-embedding-001:3072": (3072, 1536, 768),
        "gemini:gemini-embedding-2:3072": (3072, 1536, 768)}

ROWS_SQL = f"""
    SELECT t.id, t.source_id, t.title, t.body, i.category, i.year, md5(core.title_key(t.title)),
           length(core.title_key(t.title)), d.idea_id IS NOT NULL
      FROM ({embed.CANDIDATE_SQL}) t JOIN core.idea i ON i.id = t.id
      LEFT JOIN core.idea_duplicate d ON d.idea_id = t.id
     ORDER BY t.id   -- 표본이 실행마다 같아야 캐시가 맞는다. 평가 표본 모집단 = 후보(원문 50자+), 2026-10-13 평가 그대로
"""


def truncate(mat, dims: int):
    import numpy as np

    m = mat[:, :dims].astype(np.float32, copy=True)
    m /= np.maximum(np.linalg.norm(m, axis=1, keepdims=True), 1e-12)
    return m


def rank_metrics(ranks: list[int]) -> dict[str, float]:
    """ranks = 1부터. 정답 없음은 넣지 않는다."""
    n = len(ranks)
    if not n:
        return {"n": 0, "r1": float("nan"), "r10": float("nan"), "mrr": float("nan")}
    return {"n": n, "r1": sum(r == 1 for r in ranks) / n, "r10": sum(r <= 10 for r in ranks) / n,
            "mrr": sum(1 / r for r in ranks) / n}


def first_hit_ranks(scores, positives: list[set[int]], exclude: list[set[int]]) -> list[int]:
    """scores = 질의×후보. 질의마다 제외 후보를 뺀 순위에서 첫 정답의 순위(1부터)."""
    import numpy as np

    out = []
    for q, pos in enumerate(positives):
        if not pos:
            continue
        s = scores[q].copy()
        for j in exclude[q]:
            s[j] = -np.inf
        best = max(s[j] for j in pos)
        out.append(int((s > best).sum()) + 1)
    return out


def purity(mat, cats: list[str], k: int = 10) -> tuple[float, float]:
    """(kNN@k 같은 category 비율 평균, 우연 기준선 Σp²). 자기 자신은 이웃에서 뺀다."""
    import numpy as np

    sims = mat @ mat.T
    np.fill_diagonal(sims, -np.inf)
    nn = np.argpartition(-sims, k, axis=1)[:, :k]
    arr = np.array(cats)
    pur = float((arr[nn] == arr[:, None]).mean())
    n = len(cats)
    base = sum((c / n) ** 2 for c in Counter(cats).values())
    return pur, base


def sample(rows, rng) -> tuple[list[int], set[str]]:
    """표본 E 행 번호. 대표 기록만, 소스 비례 + C 대상 소스 최소 C_MIN_ROWS."""
    by_src: dict[str, list[int]] = defaultdict(list)
    for n, r in enumerate(rows):
        if not r[8]:
            by_src[r[1]].append(n)
    cats: dict[str, Counter] = defaultdict(Counter)
    for r in rows:
        if r[4] and not r[8]:
            cats[r[1]][r[4]] += 1
    c_src = {s for s, c in cats.items() if sum(c.values()) >= C_MIN_ROWS and len(c) >= 2}
    total = sum(len(v) for v in by_src.values())
    picked: list[int] = []
    for s in sorted(by_src):
        idx = by_src[s]
        want = round(E_SIZE * len(idx) / total)
        if s in c_src:
            want = max(want, min(C_MIN_ROWS, len(idx)))
        picked += [int(i) for i in rng.choice(idx, size=min(want, len(idx)), replace=False)]
    return sorted(picked), c_src


def evaluate(dsn: str, out: str, *, eg: Egress | None = None, models: tuple[str, ...] = tuple(DIMS)) -> dict:
    import numpy as np

    rng = np.random.default_rng(SEED)
    own = eg is None
    eg = eg or Egress.from_dsn(dsn)
    outdir = Path(out)
    outdir.mkdir(parents=True, exist_ok=True)
    try:
        with db.connect(dsn) as conn:
            conn.execute("SET statement_timeout = '10min'")
            rows = conn.execute(ROWS_SQL).fetchall()
            plan = _plan(rows, rng)
            results, review = [], {}
            for key in models:
                results.append(_score_model(conn, eg, key, rows, plan))
                review[key] = _review(conn, key, rows, plan)
    finally:
        if own:
            eg.close()
    _write(outdir, rows, plan, results, review)
    return {"models": len(results), "rows": sum(len(r["dims"]) for r in results), "out": str(outdir),
            "calls": sum(r["calls"] for r in results)}


def _plan(rows, rng) -> dict:
    """모델과 무관한 표본·질의·정답 집합."""
    e_idx, c_src = sample(rows, rng)
    notitle_ok = {n for n in range(len(rows))
                  if embed.body_key_len(embed.input_text(_NT, rows[n][2], rows[n][3])) >= 50}
    # A: E 안에서 제목이 유일하고 제목을 지운 본문이 50자 이상
    tcount = Counter(rows[n][6] for n in e_idx)
    a_cand = [n for n in e_idx if tcount[rows[n][6]] == 1 and n in notitle_ok and rows[n][7] >= 2]
    a_q = sorted(int(i) for i in rng.choice(a_cand, size=min(A_QUERIES, len(a_cand)), replace=False))
    # B: 같은 제목(10자+) 묶음, 묶음 크기 2..B_MAX_GROUP, 묶음 안 본문이 두 가지 이상
    groups: dict[str, list[int]] = defaultdict(list)
    for n, r in enumerate(rows):
        if r[7] >= 10 and n in notitle_ok:
            groups[r[6]].append(n)
    groups = {g: m for g, m in groups.items() if 2 <= len(m) <= B_MAX_GROUP
              and len({embed.sha(embed.clean(rows[i][3])) for i in m}) > 1}
    b_pool_q = sorted(n for m in groups.values() for n in m)
    b_q = sorted(int(i) for i in rng.choice(b_pool_q, size=min(B_QUERIES, len(b_pool_q)), replace=False))
    b_members = sorted({n for q in b_q for n in groups[rows[q][6]]})
    docs = sorted((set(e_idx) & notitle_ok) | set(b_members))
    # 사람 확인 질의: E 밖 대표 기록, 소스 층화
    e_set = set(e_idx)
    by_src: dict[str, list[int]] = defaultdict(list)
    for n, r in enumerate(rows):
        if not r[8] and n not in e_set and r[7] >= 4:
            by_src[r[1]].append(n)
    big = sorted(by_src, key=lambda s: -len(by_src[s]))
    review_q: list[int] = []
    while len(review_q) < REVIEW_QUERIES and any(by_src[s] for s in big):
        for s in big:
            left = [i for i in by_src[s] if i not in review_q]
            if left and len(review_q) < REVIEW_QUERIES:
                review_q.append(int(rng.choice(left)))
            by_src[s] = left
    return {"e": e_idx, "c_src": c_src, "a_q": a_q, "a_cand": len(a_cand), "a_dropped_title_clash":
            sum(1 for n in e_idx if tcount[rows[n][6]] > 1), "docs": docs, "b_q": b_q, "b_groups": len(groups),
            "review_q": review_q}


def _rows_for(rows, idxs):
    return [(rows[n][0], rows[n][1], rows[n][2], rows[n][3]) for n in idxs]


def _score_model(conn, eg, key: str, rows, plan) -> dict:
    t0 = time.monotonic()
    calls = chars = 0
    for variant, idxs in ((key, plan["e"]), (f"{key}:notitle", plan["docs"]),
                          (f"{key}:title", sorted(set(plan["a_q"]) | set(plan["b_q"]) | set(plan["review_q"])))):
        r = embed.ensure(conn, eg, variant, _rows_for(rows, idxs))
        calls += r["calls"]
        chars += r["chars"]
    elapsed = time.monotonic() - t0
    ids = [r[0] for r in rows]
    _, full = embed.load(conn, key, [ids[n] for n in plan["e"]])
    _, docs = embed.load(conn, f"{key}:notitle", [ids[n] for n in plan["docs"]])
    _, aq = embed.load(conn, f"{key}:title", [ids[n] for n in plan["a_q"]])
    _, bq = embed.load(conn, f"{key}:notitle", [ids[n] for n in plan["b_q"]])
    tokens = conn.execute(
        """SELECT coalesce(sum(tokens_in), 0) FROM core.egress_call WHERE id IN
             (SELECT DISTINCT egress_call_id FROM core.idea_embedding WHERE model LIKE %s)""",
        (key + "%",)).fetchone()[0]
    doc_pos = {n: j for j, n in enumerate(plan["docs"])}
    a_pos = [{doc_pos[n]} if n in doc_pos else set() for n in plan["a_q"]]
    a_ex = [set() for _ in plan["a_q"]]
    b_pos, b_ex = [], []
    for q in plan["b_q"]:
        grp = rows[q][6]
        members = {doc_pos[n] for n in plan["docs"] if rows[n][6] == grp}
        b_pos.append(members - {doc_pos[q]})
        b_ex.append({doc_pos[q]})
    c_rows = [(j, rows[n][1], rows[n][4]) for j, n in enumerate(plan["e"]) if rows[n][1] in plan["c_src"] and rows[n][4]]
    by_dims = []
    for d in DIMS[key]:
        D, F = truncate(docs, d), truncate(full, d)
        a = rank_metrics(first_hit_ranks(truncate(aq, d) @ D.T, a_pos, a_ex))
        b = rank_metrics(first_hit_ranks(truncate(bq, d) @ D.T, b_pos, b_ex))
        per_src = {}
        for s in sorted(plan["c_src"]):
            js = [j for j, src, _ in c_rows if src == s]
            if len(js) > 10:
                pur, base = purity(F[js], [c for j2, src, c in c_rows if src == s])
                per_src[s] = {"n": len(js), "purity": pur, "base": base, "lift": pur / base if base else float("nan")}
        c = {"sources": len(per_src),
             "purity": sum(v["purity"] for v in per_src.values()) / max(1, len(per_src)),
             "base": sum(v["base"] for v in per_src.values()) / max(1, len(per_src)),
             "lift": sum(v["lift"] for v in per_src.values()) / max(1, len(per_src)), "per_source": per_src}
        by_dims.append({"dims": d, "a": a, "b": b, "c": c})
    truncated = sum(1 for n in plan["e"] if len(embed.prep(rows[n][2], rows[n][3])) >= embed.MAX_CHARS)
    return {"key": key, "dims": by_dims, "calls": calls, "chars": chars, "tokens": tokens,
            "seconds": round(elapsed, 1), "truncated": truncated}


def _review(conn, key: str, rows, plan, k: int = 10) -> list[dict]:
    ids = [r[0] for r in rows]
    _, q = embed.load(conn, f"{key}:title", [ids[n] for n in plan["review_q"]])
    _, full = embed.load(conn, key, [ids[n] for n in plan["e"]])
    sims = q @ full.T
    out = []
    for i, n in enumerate(plan["review_q"]):
        top = (-sims[i]).argsort()[:k]
        out.append({"query": rows[n][2], "query_id": rows[n][0], "query_source": rows[n][1],
                    "hits": [{"score": float(sims[i, j]), "id": rows[plan["e"][j]][0],
                              "title": rows[plan["e"][j]][2], "body": embed.clean(rows[plan["e"][j]][3])[:80],
                              "year": rows[plan["e"][j]][5], "source": rows[plan["e"][j]][1]} for j in top]})
    return out


def _fmt(x) -> str:
    return "-" if math.isnan(x) else f"{x:.3f}"


def _write(outdir: Path, rows, plan, results, review) -> None:
    head = ["key", "dims", "A n", "A R@1", "A R@10", "A MRR", "B n", "B R@1", "B R@10", "B MRR",
            "C sources", "C purity", "C Σp²", "C lift", "tokens(누적)", "calls(이번 실행)", "seconds(이번 실행)",
            "truncated"]
    table = []
    for r in results:
        for d in r["dims"]:
            a, b, c = d["a"], d["b"], d["c"]
            table.append([r["key"], d["dims"], a["n"], a["r1"], a["r10"], a["mrr"], b["n"], b["r1"], b["r10"],
                          b["mrr"], c["sources"], c["purity"], c["base"], c["lift"], r["tokens"], r["calls"],
                          r["seconds"], r["truncated"]])
    n_target = len(rows)
    lines = [
        "# 임베딩 후보 평가 (AC1)", "",
        f"- 대상: {n_target}건(비퇴역·본문 50자+), 대표 기록 {sum(1 for r in rows if not r[8])}건",
        f"- 표본 E {len(plan['e'])}건, seed {SEED}, PREP_VERSION {embed.PREP_VERSION}, MAX_CHARS {embed.MAX_CHARS}",
        (f"- A 질의 {len(plan['a_q'])}건(후보 {plan['a_cand']}, E 안 제목 충돌로 뺀 행 {plan['a_dropped_title_clash']}),"
         f" 문서 {len(plan['docs'])}건(:notitle)"),
        (f"- B 질의 {len(plan['b_q'])}건(같은 제목 묶음 {plan['b_groups']}개). 계획의 '같은 공모·다른 소스' 쌍은 실측 0쌍이라"
         " 같은 사업명(제목 10자+) 이웃 찾기로 대신함 — 주제 근접도 시험"),
        f"- C 소스 {len(plan['c_src'])}개: {', '.join(sorted(plan['c_src']))}",
        "- 차원 3072 미만 행 = 3072 벡터 앞부분을 잘라 재정규화(Matryoshka)", "",
        "| " + " | ".join(head) + " |", "|" + "---|" * len(head),
    ]
    for t in table:
        lines.append("| " + " | ".join(_fmt(x) if isinstance(x, float) else str(x) for x in t) + " |")
    lines += ["", "## 결정", "", "(사용자 결정 후 기록: 날짜, 모델 키·차원, 요지)", ""]
    (outdir / "embed-eval.md").write_text("\n".join(lines), encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(head)
    for t in table:
        ws.append(t)
    ws2 = wb.create_sheet("c_per_source")
    ws2.append(["key", "dims", "source", "n", "purity", "base", "lift"])
    for r in results:
        for d in r["dims"]:
            for s, v in d["c"]["per_source"].items():
                ws2.append([r["key"], d["dims"], s, v["n"], v["purity"], v["base"], v["lift"]])
    wb.save(outdir / "embed-eval.xlsx")

    wb = Workbook()
    wb.remove(wb.active)
    md = ["# 검색 결과 비교 (모델별 top-10, 표본 E 안에서 검색)", ""]
    keys = list(review)
    for qi, first in enumerate(review[keys[0]]):
        ws = wb.create_sheet(f"q{qi + 1:02d}")
        ws.append(["질의", first["query"], first["query_source"]])
        md += [f"## q{qi + 1:02d}. {first['query']} ({first['query_source']})", ""]
        for key in keys:
            ws.append([])
            ws.append([key])
            ws.append(["rank", "score", "title", "body", "year", "source", "id"])
            md += [f"**{key}**", ""]
            for rank, h in enumerate(review[key][qi]["hits"], 1):
                ws.append([rank, round(h["score"], 3), h["title"], h["body"], h["year"], h["source"], h["id"]])
                md.append(f"{rank}. {h['title']} — {h['body']} ({h['year']}, {h['source']})")
            md.append("")
    wb.save(outdir / "review.xlsx")
    (outdir / "review.md").write_text("\n".join(md), encoding="utf-8")
