"""DB 통합 테스트. 실행 조건 = BB_TEST_DSN(빈 pgvector DB, superuser).

    docker run -d --rm --name bb-testdb -e POSTGRES_PASSWORD=t -p 55432:5432 pgvector/pgvector:pg16
    BB_TEST_DSN=postgresql://postgres:t@localhost:55432/postgres uv run pytest tests/test_db_integration.py
"""

import os
from datetime import date

import psycopg
import pytest

from bluebird import db, export, wording
from bluebird.db import MigrationChecksumError
from bluebird.signals import catalog

DSN = os.environ.get("BB_TEST_DSN")
pytestmark = pytest.mark.skipif(not DSN, reason="BB_TEST_DSN not set")

HEADER = list(catalog.COLS)


def _csv(path, rows):
    import csv
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for r in rows:
            w.writerow([r.get(h, "") for h in HEADER])
    return path


def _row(pk, title, org_code="100", registered="2015-01-01"):
    return {"목록키": pk, "목록유형": "FILE", "목록명": title, "제공기관코드": org_code, "제공기관": "기관",
            "등록일": registered, "수정일": registered, "설명": "d", "이용허락범위": "제한 없음"}


@pytest.fixture()
def fresh(tmp_path):
    name = "bbtest"
    with psycopg.connect(DSN, autocommit=True) as c:
        c.execute(f"DROP DATABASE IF EXISTS {name} WITH (FORCE)")
        c.execute(f"CREATE DATABASE {name}")
    dsn = DSN.rsplit("/", 1)[0] + f"/{name}"
    db.migrate(dsn, "core")
    return dsn


def test_migration_checksum_guard(fresh, monkeypatch):
    files = db.migration_files("core")
    monkeypatch.setattr(db, "migration_files", lambda t: [(files[0][0], "0" * 64, files[0][2])])
    with pytest.raises(MigrationChecksumError):
        db.migrate(fresh, "core")


def test_catalog_tiers_and_reregistration(fresh, tmp_path):
    s0 = _csv(tmp_path / "s0.csv", [_row("1", "서울시_자전거 대여소_20240101"), _row("2", "기존 데이터")])
    catalog.import_snapshot(dsn=fresh, path=s0, taken_at=date(2026, 10, 6))
    s1 = _csv(tmp_path / "s1.csv", [
        _row("1", "서울시_자전거 대여소_20240101"), _row("2", "기존 데이터"),
        _row("3", "새 데이터", registered="2026-10-08"),               # 관측 신규
        _row("4", "오래된 데이터", registered="2019-03-01"),          # 다시 나타남
        _row("5", "서울시 자전거 대여소(2026년)", registered="2026-10-09"),  # 1의 재등록
        _row("6", "서울시 자전거 대여소", org_code="999", registered="2026-10-09"),  # 다른 기관 → 신규
    ])
    res = catalog.import_snapshot(dsn=fresh, path=s1, taken_at=date(2026, 10, 12))
    assert res["snapshot_id"] == 1 and res["new"] == 4
    with psycopg.connect(fresh) as c:
        got = dict(c.execute("SELECT public_data_pk, tier || coalesce(':' || rereg_of, '') FROM core.signal_dataset"))
        changes = {r[0] for r in c.execute("SELECT ref_id FROM core.condition_change")}
    assert got == {"1": "portal_registered", "2": "portal_registered", "3": "observed_new", "4": "reappeared",
                   "5": "observed_new:1", "6": "observed_new"}
    assert changes == {"3", "4", "6"}
    # 동일 파일 재적재 없음
    assert catalog.import_snapshot(dsn=fresh, path=s1, taken_at=date(2026, 10, 19)) == {"skipped": True}


def _seed_idea(c, missing_data="[]"):
    c.execute("INSERT INTO core.source (id, name, license, url, layer, public_ok, export_grade, body_present,"
              " policy_approved_by, policy_approved_at, record_kind)"
              " VALUES ('s','n','l','',2,true,'O',true,'op',now(),'contest_award')")
    c.execute("INSERT INTO core.ingest_run (source_id, status) VALUES ('s','ok')")
    c.execute("INSERT INTO core.contest (id, source_id, name, host_org, year) VALUES ('C-1','s','대회','',2019)")
    c.execute("INSERT INTO core.idea (id, source_id, contest_id, year, title, team_kind, first_ingest_id,"
              " last_ingest_id) VALUES ('ID-2019-aaaaaaaaaa','s','C-1',2019,'제목','empty',1,1)")
    c.execute("INSERT INTO core.idea_card (idea_id, card_kind, title, missing_data, extractor)"
              " VALUES ('ID-2019-aaaaaaaaaa','local_extract','제목',%s,'human')", (missing_data,))


def test_diagnosis_requires_evidence_and_missing_data(fresh):
    with psycopg.connect(fresh) as c:
        _seed_idea(c)
        c.commit()
        c.execute("INSERT INTO core.diagnosis (idea_id, \"primary\", extractor) VALUES ('ID-2019-aaaaaaaaaa','R','human')")
        with pytest.raises(psycopg.errors.RaiseException, match="requires evidence"):
            c.commit()
    with psycopg.connect(fresh) as c:
        c.execute("INSERT INTO core.evidence (kind, url) VALUES ('manual','https://x.go.kr') ")
        c.execute("INSERT INTO core.x_evidence VALUES ('diagnosis','ID-2019-aaaaaaaaaa',1)")
        c.execute("INSERT INTO core.diagnosis (idea_id, \"primary\", extractor) VALUES ('ID-2019-aaaaaaaaaa','D','human')")
        with pytest.raises(psycopg.errors.RaiseException, match="missing_data"):
            c.commit()
    with psycopg.connect(fresh) as c:  # U는 근거 없이 가능
        c.execute("INSERT INTO core.diagnosis (idea_id, \"primary\", extractor) VALUES ('ID-2019-aaaaaaaaaa','U','human')")
        c.commit()


def test_ingest_keeps_body_metadata_on_unchanged_and_failed_attempts(fresh, tmp_path):
    from bluebird import ingest

    secret = tmp_path / "secret"
    secret.write_bytes(b"s" * 32)
    config = tmp_path / "sources.toml"
    config.write_text(
        '[[source]]\nid="official"\nadapter="mapped_contest_csv"\nfile="official.csv"\n'
        'name="공식 제안"\nlicense="l"\nlayer=2\npublic_ok=false\nexport_grade="pending"\n'
        'record_kind="contest_award"\n'
        'topic="t"\nshare=false\n'
        '[source.columns]\ntitle="제안명"\nbody=["설명"]\n', encoding="utf-8")
    source = tmp_path / "official.csv"
    source.write_text("제안명,설명\n빗물 회수,지붕의 빗물을 저장하는 장치\n", encoding="utf-8")
    options = {"dsn": fresh, "config": config, "seed_dir": tmp_path, "secret_path": secret}
    first = ingest.ingest(**options)
    assert first[0]["with_body"] == 1
    assert ingest.ingest(**options)[0]["status"] == "unchanged"
    with psycopg.connect(fresh) as conn:
        assert conn.execute("SELECT body_present FROM core.source WHERE id='official'").fetchone()[0]
    source.write_text("변경된헤더,설명\n제안,잘못된 스키마\n", encoding="utf-8")
    with pytest.raises(ValueError, match="missing required headers"):
        ingest.ingest(**options)
    with psycopg.connect(fresh) as conn:
        assert conn.execute("SELECT body_present FROM core.source WHERE id='official'").fetchone()[0]
        assert conn.execute("SELECT body FROM core.idea WHERE source_id='official'").fetchone()[0] == "지붕의 빗물을 저장하는 장치"


def test_ingest_retires_vanished_rows_and_cards_follow(fresh, tmp_path):
    """원본 파일이 갱신돼 사라진 행은 retired_at, 다시 나타나면 복귀. 절반 넘게 사라지면 멈춘다."""
    import csv

    from bluebird import cards, ingest

    (tmp_path / "secret").write_bytes(b"s" * 32)
    (tmp_path / "sources.toml").write_text(
        '[[source]]\nid="d"\nadapter="design_idea_csv"\nfile="d.csv"\nname="n"\nlicense="l"\nlayer=2\n'
        'public_ok=true\nexport_grade="O"\npolicy_approved_by="t"\npolicy_approved_at="2026-10-06"\n'
        'record_kind="contest_award"\n'
        'topic="t"\nshare=false\n',
        encoding="utf-8")

    def write(rows):
        with (tmp_path / "d.csv").open("w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["등록번호", "연도", "포상", "수상자", "제목", "내용"])
            w.writerows([[str(i), "2022", "상", "", f"제목{i}", f"제목{i}의 본문. 정류장 데이터가 없어 불편하다."]
                         for i in rows])

    def run(**kw):
        return ingest.ingest(dsn=fresh, config=tmp_path / "sources.toml", seed_dir=tmp_path,
                             secret_path=tmp_path / "secret", **kw)

    def live():
        with psycopg.connect(fresh) as c:
            return c.execute("SELECT count(*) FILTER (WHERE retired_at IS NULL), count(*) FROM core.idea").fetchone()

    write(range(1, 5))
    run()
    cards.build(dsn=fresh)
    write([1, 2, 3, 9])            # 4 사라짐, 9 신규
    assert run()[0]["retired"] == 1
    assert live() == (4, 5)
    write([1, 2, 3, 4, 9])         # 4 복귀
    run()
    assert live() == (5, 5)
    write([1])                     # 4/5 사라짐 → 파일 이상으로 중단
    with pytest.raises(ingest.IngestError, match="would be retired"):
        run()
    with pytest.raises(ingest.IngestError, match="would be retired"):
        run(force=True)            # force는 재적재만, 퇴역 허용 아님
    assert live() == (5, 5)
    rep = cards.build(dsn=fresh)
    assert rep["by_kind"]["local_extract"]["with_missing_data"] == 5  # 키 없음 → full 아님


LAW_XML = ("<?xml version='1.0' encoding='UTF-8'?><법령><기본정보><법령명_한글><![CDATA[드론법]]></법령명_한글>"
           "<공포일자>20251001</공포일자><공포번호>12345</공포번호><시행일자>20260101</시행일자></기본정보></법령>")


def _law_egress(dsn):
    import httpx

    from bluebird.egress import Egress, Policy, db_recorder

    def h(req):
        assert req.url.path == "/DRF/lawService.do" and req.url.params["MST"] == "777"
        return httpx.Response(200, text=LAW_XML)

    with psycopg.connect(dsn) as c:
        pol = Policy.load(c)
    return Egress(policy=pol, recorder=db_recorder(dsn), proxy="", transport=httpx.MockTransport(h))


def test_funnel_end_to_end_human_path(fresh, tmp_path):
    """키 없는 경로: 사람 흔적·진단(근거) → 사람이 넣은 바뀐 것 → 사람 판정 → S → 승인 → 5단계."""
    from bluebird import funnel

    s0 = _csv(tmp_path / "s0.csv", [_row("1", "드론 비행 데이터")])
    catalog.import_snapshot(dsn=fresh, path=s0, taken_at=date(2026, 10, 6))
    with psycopg.connect(fresh) as c:
        _seed_idea(c)
        c.commit()
    iid = "ID-2019-aaaaaaaaaa"

    def stage():
        with psycopg.connect(fresh) as c:
            return c.execute("SELECT s0,s1,s2,s3,s4,s5 FROM core.revival_candidate WHERE idea_id=%s",
                             (iid,)).fetchone()

    assert stage() == (True, False, False, False, False, False)
    funnel.trace_pending(dsn=fresh)  # 사람 manual 없으면 pending
    assert stage()[1] is False
    funnel.trace_manual(dsn=fresh, idea_id=iid, result="none", status=None, url=None, note="검색 결과 없음", by="t")
    assert stage()[:2] == (True, True)
    with pytest.raises(ValueError, match="requires at least one"):
        funnel.diagnose_set(dsn=fresh, idea_id=iid, cause="R", rationale="규제", evidence=[], by="t")
    funnel.diagnose_set(dsn=fresh, idea_id=iid, cause="R", rationale="비행 규제",
                        evidence=[("https://www.law.go.kr/x", "비행 승인")], by="t")
    assert stage()[2] is True
    # 종류 불일치(R 아이디어 ↔ 데이터 개방) 생성 불가
    with psycopg.connect(fresh) as c:
        c.execute("INSERT INTO core.condition_change (id, kind, ref_id, occurred_at, url, title)"
                  " VALUES ('dataset:1','dataset_opened','1','2020-01-01','https://www.data.go.kr/data/1','d')")
        c.commit()
        with pytest.raises(psycopg.errors.RaiseException, match="does not fit"):
            c.execute("INSERT INTO core.change_match (change_id, idea_id) VALUES ('dataset:1',%s)", (iid,))
    # 사람이 넣는 바뀐 것은 dataset_opened 불가, 금지 문구 불가
    with pytest.raises(ValueError):
        funnel.change_add(dsn=fresh, kind="dataset_opened", url="https://x", title="t",
                          occurred_at=date(2026, 1, 1), by="t")
    with pytest.raises(ValueError, match="law-mst"):
        funnel.change_add(dsn=fresh, kind="law_effective", url="https://www.law.go.kr/lsInfoP.do?lsiSeq=1",
                          title="드론 배송 허용", occurred_at=date(2026, 1, 1), by="t")
    with pytest.raises(ValueError, match="시행일자"):  # API의 시행일자와 다르면 거부
        funnel.change_add(dsn=fresh, kind="law_effective", url="https://www.law.go.kr/lsInfoP.do?lsiSeq=1",
                          title="드론 배송 허용", occurred_at=date(2026, 1, 2), by="t", egress=_law_egress(fresh),
                          law_mst="777")
    cid = funnel.change_add(dsn=fresh, kind="law_effective", url="https://www.law.go.kr/lsInfoP.do?lsiSeq=1",
                            title="드론 배송 허용", occurred_at=date(2026, 1, 1), by="t", egress=_law_egress(fresh),
                            law_mst="777")
    with psycopg.connect(fresh) as c:
        assert c.execute("SELECT verify_status, verify_detail->>'promulgation_no', egress_call_id IS NOT NULL"
                         " FROM core.condition_change WHERE id=%s", (cid,)).fetchone() == ("api_verified", "12345", True)
    with pytest.raises(wording.WordingError):
        funnel.match_set(dsn=fresh, change_id=cid, idea_id=iid, verdict="yes", what_changed=[],
                         how_now="그때 없던 제도", by="t")
    funnel.match_set(dsn=fresh, change_id=cid, idea_id=iid, verdict="yes", what_changed=["배송 허용"],
                     how_now="이제 배송 사업 신청 가능", by="t")
    assert stage()[3] is True
    # 바뀐 것 날짜 > as_of 다음 스냅샷 → 해당 시점 기준 미통과
    with psycopg.connect(fresh) as c:
        # 스냅샷 0(10-06) 기준일 = 다음 스냅샷(10-07). 그 이후 바뀐 것 = 스냅샷 0 기준 미통과
        c.execute("UPDATE core.condition_change SET occurred_at='2026-10-07' WHERE id=%s", (cid,))
        c.execute("INSERT INTO core.catalog_snapshot VALUES (1,'2026-10-07','f','x',0,1)")
        assert c.execute("SELECT s3 FROM core.funnel_stage(0) WHERE idea_id=%s", (iid,)).fetchone()[0] is False
        assert c.execute("SELECT s3 FROM core.funnel_stage(1) WHERE idea_id=%s", (iid,)).fetchone()[0] is True
        c.rollback()
    with pytest.raises(ValueError, match="evidence URL"):
        funnel.score_set(dsn=fresh, idea_id=iid, scores={"regulation": 5, "tech": 4}, evidence={}, by="t")
    sc = funnel.score_set(dsn=fresh, idea_id=iid, scores={"regulation": 5, "tech": 4},
                          evidence={"regulation": "https://www.law.go.kr/x", "tech": "https://a.kr"}, by="t")
    assert sc.verdict == "now"
    assert stage()[4] is True
    # 재채점 → 근거 교체(이전 근거 연결 제거)
    funnel.score_set(dsn=fresh, idea_id=iid, scores={"regulation": 5, "tech": 4},
                     evidence={"regulation": "https://www.law.go.kr/x", "tech": "https://a.kr/v2"}, by="t")
    with psycopg.connect(fresh) as c:
        assert sorted(r[0] for r in c.execute(
            "SELECT e.url FROM core.x_evidence x JOIN core.evidence e ON e.id=x.evidence_id"
            " WHERE x.target_type='timeliness'")) == ["https://a.kr/v2", "https://www.law.go.kr/x"]
    # 동일 URL 두 축 근거 → 축별 근거 행 별도(공개 가드 통과). 채점자 이름 근거에 없음
    funnel.score_set(dsn=fresh, idea_id=iid, scores={"regulation": 5, "tech": 4},
                     evidence={"regulation": "https://www.law.go.kr/x", "tech": "https://www.law.go.kr/x"},
                     by="scorer-kim")
    with psycopg.connect(fresh) as c:
        rows = c.execute("SELECT e.url, e.title, e.excerpt FROM core.x_evidence x JOIN core.evidence e"
                         " ON e.id=x.evidence_id WHERE x.target_type='timeliness'").fetchall()
    assert len(rows) == 2 and all("scorer-kim" not in f"{t}{e}" for _, t, e in rows)
    funnel.approve(dsn=fresh, idea_id=iid, by="t")
    assert stage()[5] is True
    assert funnel.report(dsn=fresh)["total"]["s5"] == 1

    def live():
        with psycopg.connect(fresh) as c:
            return c.execute("SELECT count(*) FROM core.publication WHERE revoked_at IS NULL").fetchone()[0]

    assert live() == 4  # 카드·진단·S + 바뀐 것 매칭
    # 승인 뒤 진단을 고치면 승인이 철회된다(검토한 내용만 승인 유지)
    funnel.diagnose_set(dsn=fresh, idea_id=iid, cause="R", rationale="비행 규제(수정)",
                        evidence=[("https://www.law.go.kr/x", "비행 승인")], by="t2")
    assert stage()[5] is False
    # 재승인 → 5단계, 이후 반려 → 아이디어·바뀐 것 승인 전부 철회
    funnel.approve(dsn=fresh, idea_id=iid, by="t")
    assert stage()[5] is True
    funnel.reject(dsn=fresh, idea_id=iid, by="t", note="근거 부족")
    assert stage()[5] is False and live() == 0
    # 원인을 바꾸면 맞지 않는 매칭은 rejected로 빠진다
    funnel.diagnose_set(dsn=fresh, idea_id=iid, cause="T", rationale="기술", evidence=[("https://a.kr/t", "기술")],
                        by="t")
    with psycopg.connect(fresh) as c:
        assert c.execute("SELECT status FROM core.change_match WHERE idea_id=%s", (iid,)).fetchone()[0] == "rejected"
        assert c.execute("SELECT count(*) FROM core.x_evidence WHERE target_type='diagnosis' AND target_id=%s",
                         (iid,)).fetchone()[0] == 1  # 이전 원인(R)의 근거는 떨어졌다


def test_future_change_does_not_pass(fresh):
    """미래 시행일 바뀐 것 → 3단계 미통과. 시행일이 지나면 통과 → 승인 → 5단계."""
    from bluebird import funnel

    with psycopg.connect(fresh) as c:
        _seed_idea(c)
        c.execute("INSERT INTO core.catalog_snapshot VALUES (0,'2026-10-06','f','x',0,0)")
        c.commit()
    iid = "ID-2019-aaaaaaaaaa"
    funnel.trace_manual(dsn=fresh, idea_id=iid, result="none", status=None, url=None, note="없음", by="t")
    funnel.diagnose_set(dsn=fresh, idea_id=iid, cause="M", rationale="정책", evidence=[("https://a.kr/m", "정책")],
                        by="t")
    future = date(2099, 1, 1)
    cid = funnel.change_add(dsn=fresh, kind="policy_news", url="https://a.kr/p", title="정책 시행 예정",
                            occurred_at=future, by="t")
    funnel.match_set(dsn=fresh, change_id=cid, idea_id=iid, verdict="human", what_changed=["지원"], how_now="신청 가능",
                     by="t")
    with psycopg.connect(fresh) as c:
        assert c.execute("SELECT verify_status FROM core.condition_change WHERE id=%s", (cid,)).fetchone()[0] == "attested"
        assert c.execute("SELECT s3 FROM core.revival_candidate WHERE idea_id=%s", (iid,)).fetchone()[0] is False
        c.execute("UPDATE core.condition_change SET occurred_at='2026-01-01' WHERE id=%s", (cid,))
        c.commit()
        assert c.execute("SELECT s3 FROM core.revival_candidate WHERE idea_id=%s", (iid,)).fetchone()[0] is True
    funnel.score_set(dsn=fresh, idea_id=iid, scores={"policy": 4, "tech": 4},
                     evidence={"policy": "https://a.kr/p", "tech": "https://a.kr/t"}, by="t")
    funnel.approve(dsn=fresh, idea_id=iid, by="t")
    with psycopg.connect(fresh) as c:
        assert c.execute("SELECT s5 FROM core.revival_candidate WHERE idea_id=%s", (iid,)).fetchone()[0] is True


def test_match_judge_llm_path(fresh, monkeypatch):
    """키가 있으면 후보 매칭을 LLM이 판정(egress 감사 행), status는 candidate 유지 — 승인은 사람. 키가 없으면 key_required."""
    import json as _json

    import httpx

    from bluebird import funnel
    from bluebird.egress import Egress, Policy, db_recorder

    with psycopg.connect(fresh) as c:
        _seed_idea(c, missing_data='[{"name": "배차 데이터", "excerpt": "배차 데이터가 없다"}]')
        c.execute("INSERT INTO core.diagnosis (idea_id, \"primary\", extractor) VALUES ('ID-2019-aaaaaaaaaa','D','rule')")
        c.execute("INSERT INTO core.evidence (id, kind, url) VALUES (900,'body_excerpt','https://a.kr/i')")
        c.execute("INSERT INTO core.x_evidence VALUES ('diagnosis','ID-2019-aaaaaaaaaa',900)")
        c.execute("INSERT INTO core.condition_change (id, kind, ref_id, occurred_at, url, title)"
                  " VALUES ('dataset:9','dataset_opened','9','2024-01-01','https://www.data.go.kr/data/9','버스 배차 정보')")
        c.execute("INSERT INTO core.change_match (change_id, idea_id, similarity) VALUES ('dataset:9','ID-2019-aaaaaaaaaa',0.5)")
        c.commit()
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    assert funnel.match_judge(dsn=fresh)["key_required"] is True
    monkeypatch.setenv("OPENAI_API_KEY", "k")
    sent = []

    def h(req):
        sent.append(req.content.decode())
        body = {"verdict": "partial", "what_changed": ["배차 정보 개방"], "how_now": "배차 정보로 대기시간을 알린다"}
        return httpx.Response(200, json={"choices": [{"message": {"content": _json.dumps(body, ensure_ascii=False)}}]})

    with psycopg.connect(fresh) as c:
        pol = Policy.load(c)
    eg = Egress(policy=pol, recorder=db_recorder(fresh), proxy="", transport=httpx.MockTransport(h))
    assert funnel.match_judge(dsn=fresh, egress=eg)["partial"] == 1
    with psycopg.connect(fresh) as c:
        row = c.execute("SELECT llm_verdict, status, how_now FROM core.change_match").fetchone()
        assert row == ("partial", "candidate", "배차 정보로 대기시간을 알린다")
        assert c.execute("SELECT count(*) FROM core.egress_call WHERE purpose='llm'").fetchone()[0] == 1
    assert "https://www.data.go.kr" not in sent[0]  # 이벤트 URL 미전송


def test_cards_llm_path_fallback_and_full_check(fresh, tmp_path, monkeypatch):
    """P1: 성공 → full + egress_call 연결 / 전송 실패 → 규칙(local_extract) 유지, 배치 계속 / 본문 변경 → 재생성.
    DB: full ⇒ llm + egress_call_id 강제."""
    import json as _json

    import httpx

    from bluebird import cards
    from bluebird.egress import Egress, db_recorder

    with psycopg.connect(fresh) as c:
        _seed_idea(c)
        c.execute("UPDATE core.idea SET body='정류장 위치는 있다. 저상버스 배차 데이터가 공개되지 않아 기다린다.'")
        c.execute("INSERT INTO core.idea (id, source_id, contest_id, year, title, body, team_kind, first_ingest_id,"
                  " last_ingest_id) VALUES ('ID-2019-bbbbbbbbbb','s','C-1',2019,'둘째','둘째 본문입니다.','empty',1,1)")
        c.execute("DELETE FROM core.idea_card")
        c.commit()
        with pytest.raises(psycopg.errors.CheckViolation):
            c.execute("INSERT INTO core.idea_card (idea_id, card_kind, title, extractor)"
                      " VALUES ('ID-2019-aaaaaaaaaa','full','t','rule')")
    monkeypatch.setenv("OPENAI_API_KEY", "k")
    answer = {"problem": "대기", "solution": "예측", "missing_data": [
        {"name": "저상버스 배차 데이터", "excerpt": "저상버스 배차 데이터가 공개되지 않아 기다린다."}]}

    def h(req):
        if "둘째" in req.content.decode():
            raise httpx.ReadTimeout("slow")
        return httpx.Response(200, json={"choices": [{"message": {"content": _json.dumps(answer, ensure_ascii=False)}}]})

    def eg():
        with psycopg.connect(fresh) as c:
            from bluebird.egress import Policy
            pol = Policy.load(c)
        return Egress(policy=pol, recorder=db_recorder(fresh), proxy="", transport=httpx.MockTransport(h))

    cards.build(dsn=fresh, egress=eg())
    with psycopg.connect(fresh) as c:
        got = dict(c.execute("SELECT idea_id, card_kind || ':' || extractor || ':' || (egress_call_id IS NOT NULL)"
                             " FROM core.idea_card"))
        md = c.execute("SELECT missing_data FROM core.idea_card WHERE idea_id='ID-2019-aaaaaaaaaa'").fetchone()[0]
        audited = c.execute("SELECT count(*) FROM core.egress_call WHERE purpose='llm'").fetchone()[0]
    assert got == {"ID-2019-aaaaaaaaaa": "full:llm:true", "ID-2019-bbbbbbbbbb": "local_extract:rule:false"}
    assert md[0]["name"] == "저상버스 배차 데이터" and audited == 2  # 실패 호출도 감사 행 기록
    # 본문이 바뀌면 LLM 카드도 다시 만든다(본문이 비면 title_only)
    with psycopg.connect(fresh) as c:
        c.execute("UPDATE core.idea SET body=NULL WHERE id='ID-2019-aaaaaaaaaa'")
        c.commit()
    cards.build(dsn=fresh, egress=eg())
    with psycopg.connect(fresh) as c:
        assert c.execute("SELECT card_kind, extractor, missing_data FROM core.idea_card"
                         " WHERE idea_id='ID-2019-aaaaaaaaaa'").fetchone() == ("title_only", "rule", [])


def test_only_eligible_matches_are_approved(fresh):
    """적격 매칭 하나 + 미래 날짜 매칭 하나: 승인되는 것은 적격 매칭뿐."""
    from bluebird import funnel

    with psycopg.connect(fresh) as c:
        _seed_idea(c)
        c.execute("UPDATE core.source SET public_ok = true WHERE id='s'")
        c.execute("INSERT INTO core.catalog_snapshot VALUES (0,'2026-10-06','f','x',0,0)")
        c.commit()
    iid = "ID-2019-aaaaaaaaaa"
    funnel.trace_manual(dsn=fresh, idea_id=iid, result="none", status=None, url=None, note="없음", by="t")
    funnel.diagnose_set(dsn=fresh, idea_id=iid, cause="M", rationale="정책", evidence=[("https://a.kr/m", "정책")],
                        by="t")
    ok = funnel.change_add(dsn=fresh, kind="policy_news", url="https://a.kr/now", title="시행된 정책",
                           occurred_at=date(2026, 1, 1), by="t")
    later = funnel.change_add(dsn=fresh, kind="policy_news", url="https://a.kr/later", title="예정 정책",
                              occurred_at=date(2099, 1, 1), by="t")
    for cid in (ok, later):
        funnel.match_set(dsn=fresh, change_id=cid, idea_id=iid, verdict="human", what_changed=["정책"],
                         how_now="신청 가능", by="t")
    funnel.score_set(dsn=fresh, idea_id=iid, scores={"policy": 4, "tech": 4},
                     evidence={"policy": "https://a.kr/now", "tech": "https://a.kr/t"}, by="t")
    assert len(funnel.approve(dsn=fresh, idea_id=iid, by="t")["changes"]) == 1
    # 승인 후 미래 매칭 강제 approved
    with psycopg.connect(fresh) as c:
        c.execute("UPDATE core.change_match SET status='approved' WHERE change_id=%s", (later,))
        c.execute("INSERT INTO core.publication (target_type, target_id, scope, approved_by)"
                  " SELECT 'change_match', id::text, 'change', 't' FROM core.change_match WHERE change_id=%s", (later,))
        c.commit()
    # 후보에서만 쓰는 바뀐 것을 고쳐도 공개는 그대로, 승인 매칭이 쓰는 바뀐 것의 확인 상태가 바뀌면 철회
    def live():
        with psycopg.connect(fresh) as c:
            return c.execute("SELECT count(*) FROM core.publication WHERE target_type='idea' AND target_id=%s"
                             " AND revoked_at IS NULL", (iid,)).fetchone()[0]
    with psycopg.connect(fresh) as c:
        c.execute("UPDATE core.change_match SET status='candidate' WHERE change_id=%s", (later,))
        c.execute("UPDATE core.condition_change SET title='예정 정책(수정)' WHERE id=%s", (later,))
        c.commit()
    assert live() == 3
    with psycopg.connect(fresh) as c:
        c.execute("UPDATE core.condition_change SET verify_status='fetched' WHERE id=%s", (ok,))
        c.commit()
    assert live() == 0


def test_manual_announcement_path_reaches_stage3(fresh):
    """키 없는 공고 경로(G7 대체): 사람이 K-Startup URL을 넣으면 egress로 열어 fetched, 원인 C 아이디어와 매칭 → 3단계."""
    import httpx

    from bluebird import announce, funnel
    from bluebird.egress import Egress, Policy, db_recorder

    with psycopg.connect(fresh) as c:
        _seed_idea(c)
        c.execute("INSERT INTO core.catalog_snapshot VALUES (0,'2026-10-06','f','x',0,0)")
        c.commit()
        pol = Policy.load(c)
    eg = Egress(policy=pol, recorder=db_recorder(fresh), proxy="",
                transport=httpx.MockTransport(lambda req: httpx.Response(200, text="<html>공고</html>")))
    aid = announce.add(dsn=fresh, url="https://www.k-startup.go.kr/web/contents/bizpbanc-ongoing.do?pbancSn=176001",
                       title="2026 예비창업패키지", org="창업진흥원", apply_from=date(2026, 3, 1),
                       apply_to=date(2026, 3, 31), summary=None, by="t", egress=eg)
    iid = "ID-2019-aaaaaaaaaa"
    funnel.trace_manual(dsn=fresh, idea_id=iid, result="none", status=None, url=None, note="없음", by="t")
    funnel.diagnose_set(dsn=fresh, idea_id=iid, cause="C", rationale="사업화 자금", evidence=[("https://a.kr/c", "자금")],
                        by="t")
    with psycopg.connect(fresh) as c:
        st = c.execute("SELECT verify_status, egress_call_id IS NOT NULL FROM core.condition_change WHERE ref_id=%s",
                       (aid,)).fetchone()
    assert st == ("fetched", True)
    funnel.match_set(dsn=fresh, change_id=f"announcement:{aid}", idea_id=iid, verdict="human",
                     what_changed=["예비창업 지원"], how_now="예비창업패키지에 신청할 수 있다", by="t")
    with psycopg.connect(fresh) as c:
        assert c.execute("SELECT s3 FROM core.revival_candidate WHERE idea_id=%s", (iid,)).fetchone()[0] is True


def test_record_kind_audience_and_dup_are_markers_not_filters(fresh, tmp_path):
    """0011·0012: 기록 종류·학생 대회 표시·동일 제목 묶음·제외 후보 이유 = 선별 재료만. 깔때기·카드 = 전체.
    레지스트리 변경(--policy-only) → record_kind·exclude_reason 반영. 사람 입력 이유 = 재계산 생존."""
    import csv

    from bluebird import cards, ingest

    (tmp_path / "secret").write_bytes(b"s" * 32)
    cfg = tmp_path / "sources.toml"

    def write_cfg(kind, reason=None, share=False):
        cfg.write_text(
            '[[source]]\nid="m"\nadapter="mapped_contest_csv"\nfile="m.csv"\nname="n"\nlicense="l"\nlayer=2\n'
            'public_ok=true\nexport_grade="O"\npolicy_approved_by="t"\npolicy_approved_at="2026-10-06"\n'
            f'record_kind="{kind}"\n' + f'topic="t"\nshare={str(share).lower()}\n' + (f'exclude_reason="{reason}"\n' if reason else '')
            + '[source.columns]\ntitle="제목"\ncontest="대회"\nyear="연도"\n', encoding="utf-8")

    def reasons():
        with psycopg.connect(fresh) as c:
            return {t: sorted(r) for t, r in c.execute(
                "SELECT i.title, array_remove(array_agg(x.reason ORDER BY x.reason), NULL)"
                "  FROM core.idea i LEFT JOIN core.idea_exclusion x ON x.idea_id = i.id GROUP BY i.title")}

    with (tmp_path / "m.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["제목", "대회", "연도"])
        w.writerows([
            ["저상버스 도착 알림 서비스", "공공데이터 창업경진대회", "2021"],
            ["저상버스 도착알림 서비스", "공공데이터 창업경진대회", "2022"],   # 공백만 다른 중복
            ["자유낙하 실험기", "제40회 전국학생과학발명품경진대회", "2021"],
            ["캠퍼스 공유 자전거", "대학창의발명대회", "2021"],                  # '대학생'은 학생 아님
        ])

    write_cfg("contest_submission")
    ingest.ingest(dsn=fresh, config=cfg, seed_dir=tmp_path, secret_path=tmp_path / "secret")
    cards.build(dsn=fresh)
    with psycopg.connect(fresh) as c:
        assert c.execute("SELECT record_kind FROM core.source WHERE id='m'").fetchone()[0] == "contest_submission"
        assert dict(c.execute("SELECT name, audience FROM core.contest").fetchall()) == {
            "공공데이터 창업경진대회": "unknown", "제40회 전국학생과학발명품경진대회": "student", "대학창의발명대회": "unknown"}
        assert c.execute("SELECT count(*) FROM core.idea_card").fetchone()[0] == 4          # 학생 대회도 카드
        assert c.execute("SELECT count(*) FROM core.funnel_stage(-1)").fetchone()[0] == 4   # 전부 깔때기
        c.execute("UPDATE core.contest SET audience='general' WHERE name='대학창의발명대회'")  # 사람이 확정
        c.execute("INSERT INTO core.idea_exclusion (idea_id, reason, decided_by)"
                  " SELECT id, 'internal_activity', 'reviewer' FROM core.idea WHERE title='캠퍼스 공유 자전거'")
        c.commit()
    assert reasons() == {"저상버스 도착 알림 서비스": [], "저상버스 도착알림 서비스": [],
                         "자유낙하 실험기": ["student_contest"], "캠퍼스 공유 자전거": ["internal_activity"]}

    write_cfg("public_proposal", reason="certified_product", share=True)
    r = ingest.sync_policy(dsn=fresh, configs=[cfg])
    assert r["synced"] == ["m"] and r["contests_classified"] == 0 and r["exclusion_rows"] == 9   # 4건×(기록 종류+편집 판단) + 학생 1
    with psycopg.connect(fresh) as c:
        assert c.execute("SELECT record_kind, exclude_reason, share_ok FROM core.source WHERE id='m'").fetchone() == (
            "public_proposal", "certified_product", True)
        assert c.execute("SELECT audience FROM core.contest WHERE name='대학창의발명대회'").fetchone()[0] == "general"
        assert c.execute("SELECT count(*) FROM core.funnel_stage(-1)").fetchone()[0] == 4   # 이유 부착해도 제외 없음
        assert c.execute("SELECT count(*) FROM core.idea_card").fetchone()[0] == 4
    # 소스 이유 2종(기록 종류·편집 판단) = 전부 / 학생·중복 = 해당 건 / 사람 입력 = 생존
    assert reasons() == {
        "저상버스 도착 알림 서비스": ["certified_product", "public_proposal"],
        "저상버스 도착알림 서비스": ["certified_product", "public_proposal"],
        "자유낙하 실험기": ["certified_product", "public_proposal", "student_contest"],
        "캠퍼스 공유 자전거": ["certified_product", "internal_activity", "public_proposal"]}

    write_cfg("contest_award", reason="student_contest")       # 소스 단위로 쓸 수 없는 이유
    with pytest.raises(ValueError, match="exclude_reason must be one of"):
        ingest.sync_policy(dsn=fresh, configs=[cfg])

    cfg.write_text(cfg.read_text(encoding="utf-8").replace('record_kind="contest_award"\n'
        'topic="t"\nshare=false\n', ''), encoding="utf-8")
    with pytest.raises(ValueError, match="record_kind must be one of"):
        ingest.sync_policy(dsn=fresh, configs=[cfg])


def test_duplicate_record_and_thin_record_markers(fresh, tmp_path):
    """0015: 같은 본문 = 같은 아이디어(재제출), 같은 대회·연도±1·제목 + 다른 소스 = 같은 출품작.
    대회명 없는 흔한 제목은 묶지 않고 본문 없으면 thin_record. 대표 = 공개 허용 소스 우선. 표시만, 걸러내지 않음."""
    import csv

    from bluebird import ingest

    (tmp_path / "secret").write_bytes(b"s" * 32)
    cfg = tmp_path / "sources.toml"
    cfg.write_text("".join(
        f'[[source]]\nid="{sid}"\nadapter="mapped_contest_csv"\nfile="{sid}.csv"\nname="{sid}"\nlicense="l"\nlayer=2\n'
        f'public_ok={public}\nexport_grade="O"\npolicy_approved_by="t"\npolicy_approved_at="2026-10-06"\n'
        'record_kind="contest_award"\ntopic="t"\nshare=true\n'
        '[source.columns]\ntitle="제목"\ncontest="대회"\nyear="연도"\naward="상"\nbody=["본문"]\n'
        for sid, public in (("a", "true"), ("b", "false"))), encoding="utf-8")
    park = "[제안취지] 어린이 공원별로 캐릭터를 정해 디자인하면 아이들이 더 많이 찾아오고 동네 공원이 다시 살아납니다. [사업내용] 공원 캐릭터 조형물 설치"
    rows = {
        "a": [["구멍자국이 안생기는 콤파스", "제37회 전국학생과학발명품경진대회", "2015", "", ""],
              ["어린이공원도 재밌어야 간다", "시민참여예산", "2018", "", park],
              ["어린이 공원 캐릭터 디자인", "시민참여예산", "2019", "", park.replace(" ", "  ")],   # 제목 바꿔 재제출
              ["CCTV 설치", "", "2015", "", ""], ["CCTV 설치", "", "2016", "", ""]],
        "b": [["구멍자국이 안생기는 콤파스!", "제38회 전국학생과학발명품경진대회", "2016", "최우수상", "본 발명은 구멍 자국이 남지 않는 컴퍼스에 관한 것이다."],
              ["구멍자국이 안생기는 콤파스", "제39회 전국학생과학발명품경진대회", "2018", "", ""],        # 연도 2 차이
              ["CCTV 설치", "", "2015", "", ""]],                                                     # 대회명 없음
    }
    for sid, rs in rows.items():
        with (tmp_path / f"{sid}.csv").open("w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["제목", "대회", "연도", "상", "본문"])
            w.writerows(rs)
    ingest.ingest(dsn=fresh, config=cfg, seed_dir=tmp_path, secret_path=tmp_path / "secret")

    def marks():
        with psycopg.connect(fresh) as c:
            return sorted(c.execute(
                "SELECT i.source_id, i.year, i.title, ci.source_id, ci.year, d.rules,"
                "       array(SELECT x.reason FROM core.idea_exclusion x WHERE x.idea_id = i.id"
                "              AND x.reason IN ('duplicate_record', 'thin_record') ORDER BY 1)"
                "  FROM core.idea i LEFT JOIN core.idea_duplicate d ON d.idea_id = i.id"
                "  LEFT JOIN core.idea ci ON ci.id = d.canonical_id").fetchall())

    expected = [
        ("a", 2015, "CCTV 설치", None, None, None, ["thin_record"]),
        ("a", 2015, "구멍자국이 안생기는 콤파스", None, None, None, []),                 # 대표 = 공개 허용 소스
        ("a", 2016, "CCTV 설치", None, None, None, ["thin_record"]),
        ("a", 2018, "어린이공원도 재밌어야 간다", None, None, None, []),                  # 대표 = 이른 연도
        ("a", 2019, "어린이 공원 캐릭터 디자인", "a", 2018, ["same_body"], ["duplicate_record"]),
        ("b", 2015, "CCTV 설치", None, None, None, ["thin_record"]),
        ("b", 2016, "구멍자국이 안생기는 콤파스!", "a", 2015, ["same_entry"], ["duplicate_record"]),
        ("b", 2018, "구멍자국이 안생기는 콤파스", None, None, None, []),
    ]
    assert marks() == expected
    with psycopg.connect(fresh) as c:
        assert c.execute("SELECT count(*) FROM core.funnel_stage(-1)").fetchone()[0] == 8   # 표시만, 깔때기 = 전체
        c.execute("SELECT core.refresh_idea_exclusion()")                                  # 같은 트랜잭션 재실행
        c.execute("SELECT core.refresh_idea_exclusion()")
        c.commit()
    assert marks() == expected


def test_export_xlsx_marks_share_from_source_flag(fresh, tmp_path):
    """export: xlsx 1파일 전체. share 열 = core.source.share_ok(레지스트리 share). 소스 표 = 소스 수 일치."""
    from openpyxl import load_workbook

    with psycopg.connect(fresh) as c:
        _seed_idea(c)
        c.execute("UPDATE core.source SET share_ok = true WHERE id = 's'")
        c.execute("INSERT INTO core.source (id, name, license, url, layer, public_ok, export_grade, record_kind)"
                  " VALUES ('h','보류 소스','기관 약관 — 공개·외부 전송 보류','',2,false,'pending','contest_award')")
        c.execute("INSERT INTO core.contest (id, source_id, name, host_org, year) VALUES ('C-h','h','','',2020)")
        c.execute("INSERT INTO core.idea (id, source_id, contest_id, year, title, team_kind, first_ingest_id, last_ingest_id)"
                  " VALUES ('ID-2020-hhhhhhhhhh','h','C-h',2020,'보류 아이디어','empty',1,1)")
        c.commit()

    def ideas(path):
        ws = load_workbook(path, read_only=True)["ideas"]
        rows = list(ws.iter_rows(values_only=True))
        return rows[0], rows[1:]

    r = export.xlsx(dsn=fresh, out=tmp_path / "a.xlsx")
    hdr, body = ideas(tmp_path / "a.xlsx")
    assert r == {"path": str(tmp_path / "a.xlsx"), "ideas": 2, "sources": 2, "share_ok": 1}
    by_source = {b[hdr.index("source_id")]: b for b in body}
    assert by_source["s"][hdr.index("share")] == "가능" and by_source["s"][hdr.index("contest")] == "대회"
    assert by_source["h"][hdr.index("share")] == "보류" and by_source["h"][hdr.index("contest")] == "보류 소스"   # 대회명 비면 소스명
    assert by_source["s"][hdr.index("canonical_id")] is None                                   # 묶음 없음 = 빈 값


def test_master_db_every_table_column_view_has_provenance_comment():
    """정제 DB 표·칸·조회용 뷰 전부 설명(COMMENT) 있음, 형식 '[출처] …'. 새 칸 추가 시 설명 빠지면 실패."""
    name = "bbtest_comments"
    with psycopg.connect(DSN, autocommit=True) as c:
        c.execute(f"DROP DATABASE IF EXISTS {name} WITH (FORCE)")
        c.execute(f"CREATE DATABASE {name}")
    dsn = DSN.rsplit("/", 1)[0] + f"/{name}"
    db.migrate(dsn, "master")
    with psycopg.connect(dsn) as c:
        rel = c.execute("""SELECT n.nspname || '.' || c.relname, obj_description(c.oid, 'pg_class')
                             FROM pg_class c JOIN pg_namespace n ON n.oid = c.relnamespace
                            WHERE n.nspname IN ('public', 'migrations') AND c.relkind IN ('r', 'v')""").fetchall()
        col = c.execute("""SELECT n.nspname || '.' || c.relname || '.' || a.attname, col_description(c.oid, a.attnum)
                             FROM pg_class c JOIN pg_namespace n ON n.oid = c.relnamespace
                             JOIN pg_attribute a ON a.attrelid = c.oid AND a.attnum > 0 AND NOT a.attisdropped
                            WHERE n.nspname IN ('public', 'migrations') AND c.relkind = 'r'""").fetchall()
    assert {r for r, _ in rel} >= {"public.master", "public.raw", "public.quality", "public.v_usable"}
    missing = [n for n, d in rel + col if not (d and d.startswith("["))]
    assert missing == []


def test_master_egress_records_in_master_db_without_core():
    """정제 DB 명령 외부 호출 = core 없이 정제 DB egress_call에 기록. 정책 = 레지스트리 export_grade."""
    import httpx

    from bluebird.egress import Egress, EgressBlocked

    name = "bbtest_egress"
    with psycopg.connect(DSN, autocommit=True) as c:
        c.execute(f"DROP DATABASE IF EXISTS {name} WITH (FORCE)")
        c.execute(f"CREATE DATABASE {name}")
    dsn = DSN.rsplit("/", 1)[0] + f"/{name}"
    db.migrate(dsn, "master")
    transport = httpx.MockTransport(lambda r: httpx.Response(200, text="ok"))
    sources = [{"id": "a", "export_grade": "pending"}, {"id": "d", "export_grade": "deny"}]
    with Egress.for_master(dsn, sources, transport=transport) as eg:
        assert eg.get("collect", "https://www.science.go.kr/x", fields={"title": ("a", "제목")}).status == 200
        with pytest.raises(EgressBlocked):
            eg.get("collect", "https://www.science.go.kr/x", fields={"title": ("d", "제목")})
        with pytest.raises(EgressBlocked):
            eg.get("collect", "https://www.science.go.kr/x", fields={"person_name": ("a", "홍길동")})
    with psycopg.connect(dsn) as c:
        rows = c.execute("SELECT decision, source_ids, http_status FROM egress_call ORDER BY id").fetchall()
    assert rows == [("allowed", ["a"], 200), ("blocked", ["d"], None), ("blocked", ["a"], None)]


def test_provenance_groups_by_platform_and_contest(tmp_path):
    """DATA.md: 수집처·대회·제도별 행 = master 중복 제외 실제 건수. KIPRIS = 행의 대회명으로 나눔."""
    name = "bbtest_master"
    with psycopg.connect(DSN, autocommit=True) as c:
        c.execute(f"DROP DATABASE IF EXISTS {name} WITH (FORCE)")
        c.execute(f"CREATE DATABASE {name}")
    dsn = DSN.rsplit("/", 1)[0] + f"/{name}"
    db.migrate(dsn, "master")
    with psycopg.connect(dsn) as c:
        c.execute("INSERT INTO raw (source, key, payload, fetched) VALUES"
                  " ('kipris_idea_api','1','{}','2026-10-14'), ('kipris_idea_api','2','{}','2026-10-14'),"
                  " ('kipris_idea_api','3','{}','2026-10-14'), ('seoul_budget_proposals','1','{}','2026-10-12'),"
                  " ('seoul_budget_proposals.detail','1','{}','2026-10-12'), ('seoul_budget_selected','1','{}','2026-10-12')")
        c.execute("INSERT INTO master (api, key, title, body, contest, origin, url, year, fetched, award) VALUES"
                  " ('kipris_idea_api','1','가','본문','제69회 전국과학전람회','O','u',2020,now(),'대상'),"
                  " ('kipris_idea_api','2','가','본문','전국과학전람회','O','u',2021,now(),'대상'),"
                  " ('kipris_idea_api','3','나',NULL,'2015 대한민국학생발명전시회','O','u',2015,now(),NULL),"
                  " ('seoul_budget_proposals','1','다','본문',NULL,NULL,'u',2019,now(),NULL)")
        c.execute("UPDATE master SET exclude_reasons = '{duplicate}', canonical_api = 'kipris_idea_api',"
                  " canonical_key = '1' WHERE key = '2'")
    reg = tmp_path / "r.toml"
    entry = ('[[source]]\nid = "{id}"\nadapter = "awards_csv"\nfile = "x.csv"\nname = "{id}"\nlicense = "x"\n'
             'layer = 1\npublic_ok = true\nrecord_kind = "{kind}"\ntopic = "t"\nshare = true\ncollector = "{col}"\n'
             'url = "{url}"\n')
    reg.write_text(entry.format(id="kipris_idea_api", kind="contest_award", col="kipris_idea_db",
                                url="https://plus.kipris.or.kr/x")
                   + entry.format(id="seoul_budget_proposals", kind="public_proposal", col="seoul_budget",
                                  url="https://data.seoul.go.kr/x"), encoding="utf-8")
    from bluebird import master, programs

    assert programs.classify(dsn=dsn, sources=master.registry([reg])) == {
        "status": "ok", "rows": 4, "programs": 3, "removed_programs": 0, "unclassified_rows": 0}
    out = tmp_path / "DATA.md"
    assert export.provenance(dsn=dsn, configs=[reg], out=out) == {"path": str(out), "platforms": 2, "programs": 3,
                                                                  "rows": 3}
    md = out.read_text(encoding="utf-8")
    assert "- 기준: 2026-10-14" in md and "원본(raw) 6행 = 목록 5 + 상세 페이지 본문 1" in md
    assert "정제(master) 4행 = 중복 제외 3 + 중복 1" in md and "## 수집처 2곳 · 데이터셋 2개" in md
    assert "| KIPRIS Plus | API 1 | 0 | 2 |" in md and "| 서울 열린데이터광장 | API 1 | 1 | 1 |" in md
    assert "| 공모전 | 과학·발명 | 2 | 2 | 1 | 1 |" in md and "| 시민 제안 | 주민참여예산 | 1 | 1 | 1 | 0 |" in md
    assert "| 전국과학전람회 | 2020 | 1 | 1 | 1 | KIPRIS Plus |" in md
    assert "| 대한민국학생발명전시회 | 2015 | 1 | 0 | 0 | KIPRIS Plus |" in md
    assert "- `seoul_budget_selected` 1행" in md


def test_db_log_history_rules_and_merge(tmp_path, monkeypatch):
    """db_log: 이전 이력 7행 + 적는 규칙(0009). 추가만, dump = 파일 이름, restore = sha256 필수, 다른 DB 기록 옮김."""
    dsns = []
    for name in ("bbtest_dblog", "bbtest_dblog_old"):
        with psycopg.connect(DSN, autocommit=True) as c:
            c.execute(f"DROP DATABASE IF EXISTS {name} WITH (FORCE)")
            c.execute(f"CREATE DATABASE {name}")
        dsns.append(DSN.rsplit("/", 1)[0] + f"/{name}")
    dsn, old = dsns
    monkeypatch.setenv("BB_LOG_PLACE", "test")
    from bluebird import cli, dblog

    assert cli.main(["migrate", "--target", "master", "--dsn", dsn]) == 0
    rows = dblog.recent(dsn=dsn)
    assert [r[3] for r in rows] == ["files", "dump", "restore", "migrate", "dump", "restore", "drop", "migrate"]
    assert rows[-1][2] == "test" and "0008_db_log.sql, 0009_db_log_rules.sql" in rows[-1][7]
    with psycopg.connect(dsn) as c:
        c.execute("INSERT INTO raw (source, key, payload, fetched) VALUES ('s','1','{}',now()), ('s','2','{}',now())")
    dump = tmp_path / "x.dump.xz"
    dump.write_bytes(b"abc")
    sha = "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"
    assert cli.main(["master", "log", "--dsn", dsn, "--action", "dump", "--note", "보냄",
                     "--file", str(tmp_path / "아직없음.dump.xz")]) == 0                     # 덤프 전 = 이름만
    assert cli.main(["master", "log", "--dsn", dsn, "--action", "restore", "--note", "받음", "--file", str(dump)]) == 0
    assert cli.main(["master", "reasons", "--dsn", dsn]) == 0
    with psycopg.connect(dsn) as c:
        last = c.execute("SELECT place, action, note, master, raw, file, sha256 FROM db_log ORDER BY id DESC LIMIT 3"
                         ).fetchall()
    assert last[2] == ("test", "dump", "보냄", 0, 2, "아직없음.dump.xz", None)
    assert last[1] == ("test", "restore", "받음", 0, 2, "x.dump.xz", sha)
    assert last[0][:2] == ("test", "reasons") and last[0][3:5] == (0, 2)
    for sql in ("UPDATE db_log SET note = 'x' WHERE id = 1", "DELETE FROM db_log WHERE id = 1", "TRUNCATE db_log",
                "INSERT INTO db_log (place, action, note) VALUES ('test', 'restore', '파일 없음')",
                "INSERT INTO db_log (place, action, note, file) VALUES ('test', 'restore', 'sha 없음', 'f')",
                "INSERT INTO db_log (place, action, note) VALUES ('test', 'oops', 'x')",
                "INSERT INTO db_log (place, action, note) VALUES ('Server 1', 'note', 'x')",
                "INSERT INTO db_log (place, action, note) VALUES ('test', 'note', ' ')"):
        with psycopg.connect(dsn) as c, pytest.raises(psycopg.errors.Error):
            c.execute(sql)
    with pytest.raises(SystemExit):                                                   # restore = --file 필수
        cli.main(["master", "log", "--dsn", dsn, "--action", "restore", "--note", "x"])
    assert cli.main(["migrate", "--target", "master", "--dsn", dsn]) == 0            # 적용 없음 = 기록 없음
    assert len(dblog.recent(dsn=dsn, limit=100)) == 11
    # 받는 쪽: 기존 DB에만 있던 기록 → 새 DB. 같은 이력은 건너뜀
    db.migrate(old, "master")
    dblog.write(dsn=old, action="note", note="기존 DB에서만 한 일", where="server", at="2026-10-18 10:00+09")
    assert dblog.merge(dsn=dsn, src_dsn=old) == {"status": "ok", "source_rows": 8, "added": 1}
    assert dblog.merge(dsn=dsn, src_dsn=old)["added"] == 0
    merged = [r for r in dblog.recent(dsn=dsn, limit=100) if r[2] == "server" and r[3] == "note"]
    assert [(r[1], r[7]) for r in merged] == [("2026-10-18 10:00", "기존 DB에서만 한 일")]
