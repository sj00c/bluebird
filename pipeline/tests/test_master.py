import json

import httpx
import openpyxl
import pytest

from bluebird import master
from bluebird.egress import Egress, Policy


def test_read_file_cp949_rows_numbered(tmp_path):
    path = tmp_path / "geoje.csv"
    path.write_bytes("회계연도,사업명,사업비\n2017,가로등 설치,15000\n,,\n2018,쉼터,\n".encode("cp949"))
    assert master.read_file(path) == [
        ("1", {"회계연도": "2017", "사업명": "가로등 설치", "사업비": "15000"}),
        ("3", {"회계연도": "2018", "사업명": "쉼터", "사업비": ""}),
    ]


def test_read_file_utf8_bom_key_column_and_short_rows(tmp_path):
    path = tmp_path / "crawl.csv"
    path.write_bytes("\ufeffkey,title,year\n424:41975,세면대,2015\n424:41976,우산\n".encode())
    assert master.read_file(path) == [
        ("424:41975", {"key": "424:41975", "title": "세면대", "year": "2015"}),
        ("424:41976", {"key": "424:41976", "title": "우산", "year": ""}),
    ]


@pytest.mark.parametrize("text, error", [
    ("key,title\n1,a\n1,b\n", "duplicate key 1"),
    ("a,a\n1,2\n", "duplicate header"),
    ("a,b\n1,2,3\n", "row 1 has 3 cells"),
])
def test_read_file_rejects_ambiguous_rows(tmp_path, text, error):
    path = tmp_path / "bad.csv"
    path.write_text(text, encoding="utf-8")
    with pytest.raises(ValueError, match=error):
        master.read_file(path)


def test_read_file_xlsx_first_sheet(tmp_path):
    book = openpyxl.Workbook()
    book.active.append(["No.", "수상연도", "아이템명"])
    book.active.append([1, 2025, "Safety365"])
    book.save(tmp_path / "awards.xlsx")
    assert master.read_file(tmp_path / "awards.xlsx") == [("1", {"No.": "1", "수상연도": "2025", "아이템명": "Safety365"})]


def _acrc_egress(calls):
    listing = {"resultCode": "S00", "resultCount": 3, "resultList": [
        {"petiNo": "A", "title": "old"}, {"petiNo": "B", "title": "가"}, {"petiNo": "C", "title": "나"}]}

    def handler(request):
        calls.append(request.url.path.rsplit("/", 1)[1] + ":" + request.url.params.get("petiNo", ""))
        if request.url.path.endswith("OpenProposalList"):
            return httpx.Response(200, json=listing)
        if request.url.params["petiNo"] == "B":
            return httpx.Response(200, json={"resultCode": "S00", "resultData": {"petiNo": "B", "content": "본문"}})
        return httpx.Response(200, json={"resultCode": "E03", "resultMessage": "조회된 데이터가 없습니다."})

    return Egress(policy=Policy(grades={}, rules={}), recorder=lambda call: None, proxy="",
                  transport=httpx.MockTransport(handler))


def test_acrc_skips_known_falls_back_to_listing_and_stops_at_budget(monkeypatch):
    monkeypatch.setenv("DATA_GO_KR_KEY", "k")
    calls: list[str] = []
    with _acrc_egress(calls) as eg:
        rows = list(master._acrc(eg, have={"A"}, budget=3, years=range(2026, 2027)))
    assert rows == [("B", {"petiNo": "B", "content": "본문"}), ("C", {"petiNo": "C", "title": "나"})]
    assert calls == ["OpenProposalList:", "OpenProposalItem:B", "OpenProposalItem:C"]

    calls.clear()
    with _acrc_egress(calls) as eg:
        rows = list(master._acrc(eg, have={"A"}, budget=2, years=range(2026, 2027)))
    assert [k for k, _ in rows] == ["B"]


def test_gg_key_ignores_collection_date(monkeypatch):
    monkeypatch.setenv("GG_OPENAPI_KEY", "k")
    row = {"NO": "0", "BIZ_NM": "쉼터", "RCPT_YR": "2020"}

    def handler(request):
        day = "20260904" if request.url.params["pIndex"] == "1" else "20261014"
        return httpx.Response(200, json={"tbggrpbudgresdntprpbizl": [
            {"head": [{"list_total_count": 2}, {"RESULT": {"CODE": "INFO-000"}}]},
            {"row": [{**row, "DATA_COLCT_DE": day}]}]})

    with Egress(policy=Policy(grades={}, rules={}), recorder=lambda call: None, proxy="",
                transport=httpx.MockTransport(handler)) as eg:
        (k1, p1), (k2, p2) = list(master._gg(eg))
    assert k1 == k2 and p1["DATA_COLCT_DE"] != p2["DATA_COLCT_DE"]


def test_kipris_retries_html_error_page(monkeypatch):
    monkeypatch.setenv("KIPRIS_PLUS_KEY", "k")
    monkeypatch.setattr(master.time, "sleep", lambda s: None)
    replies = iter([
        httpx.Response(200, text="<html><body><p>error</body></html>"),
        httpx.Response(200, text="<response><totalCount>1</totalCount><ideaDBInfo><ideaName>우산</ideaName>"
                                 "<drawingPath>x?CN=1234567</drawingPath></ideaDBInfo></response>"),
    ])
    with Egress(policy=Policy(grades={}, rules={}), recorder=lambda call: None, proxy="",
                transport=httpx.MockTransport(lambda request: next(replies))) as eg:
        assert list(master._kipris(eg)) == [("1234567", {"ideaName": "우산", "drawingPath": "x?CN=1234567"})]


def test_registry_rejects_duplicate_ids(tmp_path):
    a, b = tmp_path / "a.toml", tmp_path / "b.toml"
    for path in (a, b):
        path.write_text('[[source]]\nid = "x"\nfile = "x.csv"\n', encoding="utf-8")
    with pytest.raises(ValueError, match="duplicate source id x"):
        master.registry([a, b])
    assert json.dumps(master.registry([a])) == '[{"id": "x", "file": "x.csv"}]'


def test_detail_body_parsers():
    seoul = ('<div class="citizen-idea"><p>시민의견 : 1</p></div><div class="txt-block">\n 최근 반려견 사고.'
             '<div>공간을 만들어 주세요.&nbsp;</div></div><div class="reply">답변</div>')
    host, _, body_of = master.DETAIL["seoul_free_suggestions"]
    assert host == "idea.seoul.go.kr"
    assert " ".join(body_of(seoul).split()) == "최근 반려견 사고. 공간을 만들어 주세요."
    kidp = ('<th scope="row" rowspan="2">상품 주요특징</th>\n<td>한글</td>\n<td>산불 &amp; 조기 진화<br>장치</td>'
            '<tr><td>영문</td><td>Wildfire</td></tr>')
    assert master.DETAIL["kidp_youth_design_projects"][2](kidp) == "산불 & 조기 진화\n장치"
    assert master.DETAIL["kidp_good_design_projects"][2]("<html>no table</html>") == ""


def test_egress_without_cookies_sends_no_session():
    seen = []

    def handler(request):
        seen.append(request.headers.get("cookie"))
        return httpx.Response(200, text="ok", headers={"set-cookie": "JSESSIONID=abc; Path=/"})

    with Egress(policy=Policy(grades={}, rules={}), recorder=lambda call: None, proxy="", cookies=False,
                transport=httpx.MockTransport(handler)) as eg:
        eg.get("collect", "https://idea.seoul.go.kr/a")
        eg.get("collect", "https://idea.seoul.go.kr/a")
    assert seen == [None, None]


def test_master_xlsx_cells_stay_text(tmp_path):
    from openpyxl import Workbook

    from bluebird.export import XLSX_CELL_MAX, _text_cell

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("master")
    ws.append([_text_cell(ws, v) for v in ("== 집값", "a\x07b", "x" * (XLSX_CELL_MAX + 5), 2025, None, "끝")])
    path = tmp_path / "m.xlsx"
    wb.save(path)
    cells = next(openpyxl.load_workbook(path).worksheets[0].iter_rows())
    assert [c.data_type for c in cells[:4]] == ["s", "s", "s", "n"]
    assert [c.value for c in cells[:2]] == ["== 집값", "ab"]
    assert len(cells[2].value) == XLSX_CELL_MAX
    assert cells[3].value == 2025 and cells[4].value is None and cells[5].value == "끝"
